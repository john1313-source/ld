"""BUDGET.md 내용을 수식이 살아 있는 엑셀(BUDGET.xlsx)로 만든다.

실행: python make_budget_xlsx.py
금액 단위: 만 원. 파란 글씨 = 입력값, 검은 글씨 = 수식, 노란 칸 = 금액 입력 필요.
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "BUDGET.xlsx")

A = "https://www.ajd.co.kr/contents/basic-tip/detail/"
SRC = {
    "평당": A + "%ED%98%84%EC%9E%A5_%EC%97%85%EC%9E%90%EA%B0%80_%EC%95%8C%EB%A0%A4%EB%93%9C%EB%A6%AC%EB%8A%94_%EC%95%84%ED%8C%8C%ED%8A%B8_%EB%A6%AC%EB%AA%A8%EB%8D%B8%EB%A7%81_%EB%B9%84%EC%9A%A9_%ED%8F%89%EB%8B%B9_%EA%B0%80%EA%B2%A9!-50541",
    "확장": A + "%EB%B0%9C%EC%BD%94%EB%8B%88_%ED%99%95%EC%9E%A5_%EB%B9%84%EC%9A%A9_%EC%96%BC%EB%A7%88%EC%9D%BC%EA%B9%8C_%EC%95%84%ED%8C%8C%ED%8A%B8_%ED%99%95%EC%9E%A5_%EA%B2%AC%EC%A0%81%EA%B3%BC_%EC%8B%9C%EA%B3%B5_%EC%A3%BC%EC%9D%98%EC%82%AC%ED%95%AD_%EC%B4%9D%EC%A0%95%EB%A6%AC-93092",
    "에어컨": "https://ohou.se/advices/12575",
    "욕실": A + "%ED%99%94%EC%9E%A5%EC%8B%A4%EB%A6%AC%EB%AA%A8%EB%8D%B8%EB%A7%81_%EB%B9%84%EC%9A%A9_%EC%9A%95%EC%8B%A4_%EB%A6%AC%EB%AA%A8%EB%8D%B8%EB%A7%81_%EB%B9%84%EC%9A%A9_%EC%8B%9C%EA%B3%B5%ED%83%80%EC%9E%85%EB%B3%84_%EC%98%88%EC%83%81_%EA%B2%AC%EC%A0%81-54498",
    "키친바흐": "https://store.hanssem.com/category/20009",
    "붙박이": A + "40%ED%8F%89%EB%8C%80_%EC%95%84%ED%8C%8C%ED%8A%B8_%EC%9D%B8%ED%85%8C%EB%A6%AC%EC%96%B4_%EB%B9%84%EC%9A%A9%EA%B3%BC_%EA%B2%AC%EC%A0%81,_%EB%AA%A8%EB%93%A0_%EC%84%B8%EB%B6%80_%EC%82%AC%ED%95%AD_%EC%B4%9D%EC%A0%95%EB%A6%AC-51804",
    "커튼": "https://www.clien.net/service/board/kin/17881781",
    "냉장고": "https://news.samsung.com/kr/%EC%82%BC%EC%84%B1%EC%A0%84%EC%9E%90-%EB%B9%84%EC%8A%A4%ED%8F%AC%ED%81%AC-%ED%82%A4%EC%B9%9C%ED%95%8F-%EB%A7%A5%EC%8A%A4-%EB%83%89%EC%9E%A5%EA%B3%A0-%EC%8B%A0%EC%A0%9C%ED%92%88-%EC%B6%9C%EC%8B%9C",
    "식세기": "https://keyzard.cc/views/nb/PS02JjBlZW1ta2xmbW5maGppbQ",
    "워시타워": "https://www.lge.co.kr/wash-tower/wa2525egzf",
    "TV": "https://www.samsung.com/sec/tvs/all-tvs/",
    "그레이슨": "https://www.lazboy.co.kr/127",
    "매트리스": "https://www.simmons.co.kr/products?categoryNo=848502",
    "프레임": "https://www.simmons.co.kr/products?categoryNo=848503",
    "식탁": "https://store.hanssem.com/goods/1116732",
    "일룸": "https://www.iloom.com/product/item.do?categoryNo=23",
}
HEARD = "전해 들음 (검색 요약)"
TODO = "금액 입력 필요"

# (구분, 항목, 모델 내용, 대조 제품·기준, 하한, 상한, 근거 수준, 출처 키, 비고)
ROWS = [
    ("공사", "올수리 기본 공사", "무몰딩·히든도어·포세린 타일·강마루·라인조명", "평당 100~250만 원 × 42.49평",
     "=42.49*100", "=42.49*250", HEARD, "평당", "평당 단가에 샷시·확장·에어컨이 포함된 견적이면 아래 항목과 중복"),
    ("공사", "샷시 교체", "확장부 이중창", "40평대 300~600만 원 이상", 300, 600, HEARD, "확장", ""),
    ("공사", "발코니 확장 — 거실, 안방", "창호벽·날개벽 전부 철거", "거실 100~150 + 안방 80~120", 180, 270, HEARD, "확장", ""),
    ("공사", "발코니 확장 — 자녀방 2곳, 주방 뒤", "창호벽·날개벽 전부 철거", "미확인", None, None, "—", None, TODO),
    ("공사", "시스템에어컨 6대", "거실·주방·안방 대형, 서재·자녀방 소형", "40평대 5~6대 850~1,200만 원", 850, 1200, HEARD,
     "에어컨", ""),
    ("공사", "욕실 2개", "안방욕실 반신욕조 750, 공용욕실 욕조", "1개당 230(평균)~950(프리미엄) × 2", "=230*2", "=950*2",
     HEARD, "욕실", ""),
    ("공사", "주방가구", "창측 하부장, 좌우 키큰장, 인덕션 아일랜드", "한샘 키친바흐 40평 1세트 (기기·시공 제외)", 900, 1500,
     HEARD, "키친바흐", ""),
    ("공사", "붙박이장·드레스룸", "자녀방 2, 안방 드레스룸", "방당 100~300 × 2 + 드레스룸 200~500", "=100*2+200",
     "=300*2+500", HEARD, "붙박이", ""),
    ("공사", "제작 가구 (신발장·팬트리·서재 책장·TV장)", "", "미확인", None, None, "—", None, TODO),
    ("공사", "커튼 (거실·안방 전동)", "쉬어 커튼", "30평대 거실·안방 전동 + 블라인드 약 150만 원 사례", 150, 200, HEARD,
     "커튼", ""),
    ("가전", "냉장고", "좌측 키큰장 빌트인", "삼성 비스포크 4도어 키친핏 맥스 (출고가)", 309, 369, HEARD, "냉장고", ""),
    ("가전", "김치냉장고", "우측 키큰장", "미확인", None, None, "—", None, TODO),
    ("가전", "식기세척기 + 인덕션", "아일랜드 인덕션", "삼성 비스포크 14인용 + 3구 인덕션 세트 (할인가~정가)", 236.2, 268,
     HEARD, "식세기", ""),
    ("가전", "천장형 후드, 오븐", "아일랜드 상부 후드, 키큰장 오븐", "미확인", None, None, "—", None, TODO),
    ("가전", "세탁기·건조기", "다용도실 직렬", "LG 트롬 워시타워 (W2520WHM ~ WA2525EGZF)", 208.3, 482, HEARD,
     "워시타워", ""),
    ("가전", "TV 85인치", "거실 아트월", "삼성 2026년형 85인치 (미니LED ~ 네오QLED 출시가)", 289, 369, HEARD, "TV", ""),
    ("가구", "거실 소파 — 레이지보이", "폭 약 3.2 m", "레이지보이 (사용자 지정 예산)", 1000, 1000,
     "확인됨 (사용자 지정)", None, "사용자 지정 1,000만 원. 제품·모델명 미정"),
    ("가구", "라운지체어 2개 (거실·안방 창가)", "1인 체어", "레이지보이 그레이슨 1인 리클라이너 × 2", "=260*2", "=299.5*2",
     HEARD, "그레이슨", ""),
    ("가구", "안방 킹 매트리스", "킹 침대", "시몬스 뷰티레스트 킹오브킹 (퓨전 할인가 ~ 자스민 정가)", 191.3, 569, HEARD,
     "매트리스", ""),
    ("가구", "안방 침대 프레임", "패브릭 헤드보드", "시몬스 뷰티레스트 시트러스 킹오브킹 (할인가~정가)", 408.7, 601, HEARD,
     "프레임", ""),
    ("가구", "식탁 세트", "2000×900 식탁 + 의자", "포세린 세라믹 2000 6인 세트 (한샘 루엘 ~ 세리드)", 72.8, 140, HEARD,
     "식탁", ""),
    ("가구", "자녀방 책상 2개", "창가 책상", "일룸 링키 컴팩트 책상세트 43.9만 원 × 2 (2026-03-01 최저가)", "=43.9*2",
     "=43.9*2", HEARD, "일룸", ""),
    ("가구", "자녀방 침대 2개", "싱글", "미확인", None, None, "—", None, TODO),
    ("가구", "서재 소파베드·의자, 화장대, 조명, 러그", "", "미확인", None, None, "—", None, TODO),
]

FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
YELLOW = PatternFill("solid", fgColor="FFFF00")
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
NORMAL = Font(name=FONT)
NUM = '#,##0.0;(#,##0.0);"-"'


def header(ws, row, names):
    for j, h in enumerate(names, 1):
        c = ws.cell(row, j, h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "예산표"
    ws["A1"] = "브라운스톤휘경 42평 재설계 — 예산 대조표 (단위: 만 원)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("조사일 2026-09-24. 금액은 웹 검색 결과 요약에서 가져온 값(전해 들음)이며 상세페이지를 직접 확인하지 못함. "
                "파란 글씨 = 입력값, 검은 글씨 = 수식, 노란 칸 = 금액 입력 필요.")
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 30

    r0 = 4
    header(ws, r0, ["구분", "항목", "모델 내용", "대조 제품·기준", "하한 (만 원)", "상한 (만 원)", "근거 수준", "출처", "비고"])
    for i, (grp, item, model, ref, lo, hi, ev, src, note) in enumerate(ROWS):
        r = r0 + 1 + i
        for j, v in enumerate([grp, item, model, ref, lo, hi, ev, None, note], 1):
            c = ws.cell(r, j, v)
            c.border, c.font = BORDER, NORMAL
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for j in (5, 6):
            c = ws.cell(r, j)
            c.number_format = NUM
            if c.value is None:
                c.fill = YELLOW
            elif isinstance(c.value, str) and c.value.startswith("="):
                c.font = BLACK
            else:
                c.font = BLUE
        if src:
            c = ws.cell(r, 8, "링크")
            c.hyperlink = SRC[src]
            c.font = Font(name=FONT, color="0563C1", underline="single")
    last = r0 + len(ROWS)
    for j, w in enumerate([7, 32, 34, 44, 12, 12, 20, 8, 40], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{r0 + 1}"
    ws.cell(last + 2, 1, "주의: 올수리 평당 단가에 샷시·확장·에어컨이 이미 포함된 견적도 있어 공사 합계에 중복이 있을 수 있음."
            ).font = Font(name=FONT, italic=True, size=9)

    s = wb.create_sheet("요약")
    s["A1"] = "예산 요약 (단위: 만 원)"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    header(s, 3, ["구분", "하한", "상한", "미확인 항목 수"])

    def rng(col):
        return f"예산표!${col}${r0 + 1}:${col}${last}"

    for i, g in enumerate(["공사", "가전", "가구"]):
        r = 4 + i
        s.cell(r, 1, g)
        s.cell(r, 2, f'=SUMIFS({rng("E")},{rng("A")},A{r})')
        s.cell(r, 3, f'=SUMIFS({rng("F")},{rng("A")},A{r})')
        s.cell(r, 4, f'=COUNTIFS({rng("A")},A{r},{rng("E")},"")')
    s.cell(7, 1, "확인된 항목 합계")
    s.cell(7, 2, "=SUM(B4:B6)")
    s.cell(7, 3, "=SUM(C4:C6)")
    s.cell(7, 4, "=SUM(D4:D6)")
    s.cell(8, 1, "총예산 (가구 포함)")
    s.cell(8, 2, 50000)
    s.cell(8, 3, "=B8")
    s.cell(9, 1, "남는 금액")
    s.cell(9, 2, "=B8-B7")
    s.cell(9, 3, "=C8-C7")
    s.cell(10, 1, "총예산 대비 사용률")
    s.cell(10, 2, "=IF(B8=0,0,B7/B8)")
    s.cell(10, 3, "=IF(C8=0,0,C7/C8)")
    for r in range(4, 11):
        for j in range(1, 5):
            c = s.cell(r, j)
            if j == 4 and r > 7:
                continue
            c.border = BORDER
            c.font = Font(name=FONT, bold=(j == 1 and r >= 7))
            if j in (2, 3):
                c.number_format = "0.0%" if r == 10 else NUM
    s["B8"].font = Font(name=FONT, color="0000FF", bold=True)
    s["B8"].comment = Comment("사용자 제시 총예산: 가구 포함 약 5억 원 (= 50,000만 원)", "Claude")
    s["A12"] = "하한·상한은 '예산표' 시트에서 금액이 확인된 항목만 합친 값. 노란 칸(미확인)에 금액을 넣으면 자동으로 반영됨."
    s["A13"] = "레이지보이 소파는 사용자 지정 1,000만 원. 나머지 금액은 웹 검색 요약(전해 들음)."
    for a in ("A12", "A13"):
        s[a].font = Font(name=FONT, italic=True, size=9)
    for j, w in enumerate([22, 16, 16, 16], 1):
        s.column_dimensions[get_column_letter(j)].width = w
    wb.move_sheet("요약", offset=-1)
    wb.active = 0
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
